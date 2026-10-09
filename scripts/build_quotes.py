#!/usr/bin/env python3
"""Turn the quote spreadsheet into data/quotes.json for the Quotes section.

Run from the repo root:
    python3 scripts/build_quotes.py                 # trial: only "Notion: Life Library" rows
    python3 scripts/build_quotes.py --all           # every row on the Master tab

Rows with Favorite = "Yes" rotate on the Life Library front page. They're always
included, even when they come from somewhere other than --source.

Options:
    --xlsx PATH      spreadsheet to read (default: Quote_Archive_Master.xlsx)
    --source NAME    only rows whose "Came from" equals NAME (default: Notion: Life Library)
    --all            ignore "Came from" and use every row
    --out PATH       where to write (default: data/quotes.json)

The private columns (My notes, Context, Screenshot file) are never copied.
Needs openpyxl:  python3 -m pip install openpyxl
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

# Spencer's labels, exactly as they should appear on the site (index.html has the same list)
LABELS = [
    "Personal Development", "Learning", "Career", "Habits", "Work", "Life",
    "Trading",
    "Writing", "Beauty", "Art",
    "Relationships", "Parenting", "Education", "Talent", "Character",
    "Personal Finance", "Economics", "Sociology",
    "Decision-Making", "Productivity", "Psychology", "Motivation", "Mindset", "Communication",
    "Management", "Technology", "Leadership", "Philosophy", "Business", "Science", "Society", "Politics",
    "Fitness", "Wellness",
]
# Spelling differences forgiven when matching a label (lowercase, spaces/dashes ignored)
CANON = {l.lower().replace("-", "").replace(" ", ""): l for l in LABELS}


def clean(v):
    if v is None:
        return ""
    return str(v).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default="Quote_Archive_Master.xlsx")
    ap.add_argument("--source", default="Notion: Life Library")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--out", default="data/quotes.json")
    args = ap.parse_args()

    ws = load_workbook(args.xlsx, read_only=True, data_only=True)["Master"]
    rows = ws.iter_rows(values_only=True)
    header = [clean(h) for h in next(rows)]
    col = {h: i for i, h in enumerate(header)}
    for need in ["ID", "Quote", "Author", "Came from", "Topics"]:
        if need not in col:
            sys.exit(f'Missing column "{need}" on the Master tab. Found: {header}')

    def get(row, name):
        i = col.get(name)
        return clean(row[i]) if i is not None and i < len(row) else ""

    quotes, unknown = [], {}
    for row in rows:
        text = get(row, "Quote")
        if not text:
            continue
        favorite = get(row, "Favorite").lower() in ("yes", "y", "true", "1")
        if not args.all and not favorite and get(row, "Came from") != args.source:
            continue
        labels = []
        for raw in get(row, "Topics").split(","):
            raw = raw.strip()
            if not raw:
                continue
            label = CANON.get(raw.lower().replace("-", "").replace(" ", ""))
            if label is None:
                unknown[raw] = unknown.get(raw, 0) + 1
            elif label not in labels:
                labels.append(label)
        qid = get(row, "ID")
        if qid.endswith(".0"):
            qid = qid[:-2]
        author = get(row, "Author")
        quotes.append({
            "id": qid,
            "text": text,
            "author": author,
            "authorFilter": get(row, "Author (for filtering)") or author,
            "attribution": get(row, "Attribution"),
            "source": get(row, "Source"),
            "link": get(row, "Link"),
            "labels": labels,
            "length": len(text),
            "favorite": favorite,
        })

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(quotes, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    print(f"Wrote {len(quotes)} quotes to {out}")
    print(f"  favorites (rotate on the front page): {sum(q['favorite'] for q in quotes)}")
    print(f"  with no label: {sum(not q['labels'] for q in quotes)}")
    if unknown:
        print("  Topics not in the label list (left off the site):")
        for name, n in sorted(unknown.items(), key=lambda x: -x[1]):
            print(f"    {name}: {n}")


if __name__ == "__main__":
    main()
